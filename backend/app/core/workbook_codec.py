"""Pure XLSX encoding, decoding, and workbook validation. No database access."""
from datetime import date, datetime, timezone
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import math
import secrets
import zipfile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from .domain import Invalid, KINDS, credit_replay, day, decimal, empty, loan_balance, replay

FORMAT_VERSION = 1

APP_VERSION = "1"

WORKBOOK_TYPE = "financial_planner_backup"

MAX_UPLOAD_BYTES = 64 * 1024 * 1024

MAX_WORKBOOK_BYTES = 64 * 1024 * 1024

MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024

MAX_SHEET_ROWS = 1_048_576

MAX_CELL_TEXT = 32_767

SHEETS = ("Metadata", "Planner State", "Instrument Catalog", "Portfolio History", "Stock Cache", "Stock Cache Data")

TABLES = ("planner_state", "instrument_catalog", "portfolio_history", "stock_cache")

HEADERS = {
    "Metadata": ("Section", "Name", "Value", "SHA-256", "Description"),
    "Planner State": ("Node ID", "Parent Node ID", "Position", "Key", "Node Type", "Value"),
    "Instrument Catalog": ("exchange", "symbol", "name", "asset_class", "currency", "native_exchange", "source", "active", "updated_at"),
    "Portfolio History": ("account", "date", "value_sgd", "gross_change_sgd", "adjusted_change_sgd", "external_flow_sgd", "updated_at"),
    "Stock Cache": ("Record ID", "security_id", "kind", "fetched_at", "expires_at"),
    "Stock Cache Data": ("Record ID", "Node ID", "Parent Node ID", "Position", "Key", "Node Type", "Value"),
}

DESCRIPTIONS = {
    "Planner State": "Typed tree for the complete planner_state row. Object keys and list positions are explicit. Values are not JSON blobs.",
    "Instrument Catalog": "One row per exchange-qualified catalog instrument. Timestamps are ISO 8601 text.",
    "Portfolio History": "One row per account and date. Monetary values stay exact text; timestamps are ISO 8601 text.",
    "Stock Cache": "Cache record keys and timestamps. Record ID links to Stock Cache Data.",
    "Stock Cache Data": "Typed tree for each stock_cache.data value. Values are not JSON blobs.",
}

class WorkbookError(ValueError):
    pass

def _timestamp(value):
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")

def _db_row(row):
    result = dict(row)
    for key, value in result.items():
        if isinstance(value, datetime):
            result[key] = _timestamp(value)
        elif isinstance(value, date):
            result[key] = value.isoformat()
    return result

def _cell_text(value):
    if isinstance(value, str) and len(value) > MAX_CELL_TEXT:
        raise WorkbookError(f"A cell value exceeds Excel's {MAX_CELL_TEXT:,}-character limit")
    return value

def _digest(rows):
    digest = sha256()
    for row in rows:
        digest.update(b"R")
        for value in row:
            if value is None:
                kind, data = b"N", b""
            elif isinstance(value, bool):
                kind, data = b"B", b"1" if value else b"0"
            elif isinstance(value, int):
                kind, data = b"I", str(value).encode("ascii")
            elif isinstance(value, float):
                kind, data = b"F", repr(value).encode("ascii")
            elif isinstance(value, str):
                kind, data = b"S", value.encode("utf-8")
            else:
                raise WorkbookError(f"Unsupported workbook cell type: {type(value).__name__}")
            digest.update(kind + len(data).to_bytes(8, "big") + data)
    return digest.hexdigest()

def _node_rows(value, prefix, record_id=None):
    rows, counter = [], 0

    def add(item, parent, position, key, depth=0):
        nonlocal counter
        if depth > 256:
            raise WorkbookError("Node tree exceeds the supported nesting depth")
        counter += 1
        if counter + 1 > MAX_SHEET_ROWS:
            raise WorkbookError("Node table exceeds Excel's worksheet row limit")
        node_id = f"{prefix}{counter:09d}"
        if isinstance(item, dict):
            node_type, scalar = "object", None
        elif isinstance(item, list):
            node_type, scalar = "list", None
        elif item is None:
            node_type, scalar = "null", None
        elif isinstance(item, bool):
            node_type, scalar = "boolean", "true" if item else "false"
        elif isinstance(item, int):
            node_type, scalar = "integer", str(item)
        elif isinstance(item, float):
            if not math.isfinite(item):
                raise WorkbookError("Non-finite numbers cannot be represented in Excel")
            node_type, scalar = "number", repr(item)
        elif isinstance(item, str):
            node_type, scalar = "string", _cell_text(item) if item else None
        else:
            raise WorkbookError(f"Unsupported nested value type: {type(item).__name__}")
        if record_id is None:
            rows.append((node_id, parent, position, key or None, node_type, scalar))
        else:
            rows.append((record_id, node_id, parent, position, key or None, node_type, scalar))
        if isinstance(item, dict):
            for child_position, (child_key, child) in enumerate(item.items()):
                if not isinstance(child_key, str):
                    raise WorkbookError("Object keys must be text")
                add(child, node_id, child_position, child_key, depth + 1)
        elif isinstance(item, list):
            for child_position, child in enumerate(item):
                add(child, node_id, child_position, None, depth + 1)

    add(value, None, 0, None)
    return rows

def _make_sheet_rows(snapshot):
    states = snapshot["planner_state"]
    if len(states) != 1 or states[0].get("id") != 1:
        raise WorkbookError("Database must contain exactly one planner_state row")
    result = {"Planner State": _node_rows(states[0], "P")}
    catalog = sorted(snapshot["instrument_catalog"], key=lambda row: (row["exchange"], row["symbol"]))
    result["Instrument Catalog"] = [tuple(_cell_text(row.get(key)) for key in HEADERS["Instrument Catalog"]) for row in catalog]
    history = sorted(snapshot["portfolio_history"], key=lambda row: (row["account"], row["date"]))
    result["Portfolio History"] = [tuple(_cell_text(row.get(key)) for key in HEADERS["Portfolio History"]) for row in history]
    cache = sorted(snapshot["stock_cache"], key=lambda row: (row["security_id"], row["kind"]))
    cache_rows, cache_data = [], []
    for index, row in enumerate(cache, 1):
        record_id = f"C{index:09d}"
        cache_rows.append((record_id, *(_cell_text(row.get(key)) for key in ("security_id", "kind", "fetched_at", "expires_at"))))
        cache_data.extend(_node_rows(row["data"], record_id + "-N", record_id))
    result["Stock Cache"], result["Stock Cache Data"] = cache_rows, cache_data
    for sheet, rows in result.items():
        if len(rows) + 1 > MAX_SHEET_ROWS:
            raise WorkbookError(f"{sheet} exceeds Excel's worksheet row limit")
    return result

def _metadata_rows(sheet_rows, exported_at, schema_version, sample):
    rows = [
        ("workbook", "workbook_type", WORKBOOK_TYPE, None, "Financial Planner complete application-data workbook."),
        ("workbook", "workbook_role", "template" if sample else "backup", None, "Template files contain an empty sample dataset."),
        ("workbook", "format_version", FORMAT_VERSION, None, "Workbook format version."),
        ("workbook", "app_version", APP_VERSION, None, "Application data format version."),
        ("workbook", "schema_version", schema_version, None, "planner_state schema version."),
        ("workbook", "exported_at", exported_at, None, "Export timestamp in UTC."),
        ("workbook", "sample_data", sample, None, "True marks the minimal sample workbook."),
    ]
    if sample:
        rows.append(("workbook", "sample_warning", "Synthetic sample data", None,
                     "Contains a clearly named example account and opening entry."))
    for sheet in SHEETS[1:]:
        rows.append(("sheet", sheet, len(sheet_rows[sheet]), _digest(sheet_rows[sheet]), DESCRIPTIONS[sheet]))
    return rows

def _append_row(ws, values, header=False):
    cells = []
    for value in values:
        value = _cell_text(value)
        cell = WriteOnlyCell(ws, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
            cell.number_format = "@"
        if header:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24445C")
        cells.append(cell)
    ws.append(cells)

def _build_xlsx(snapshot, sample=False, max_workbook_bytes=None):
    sheet_rows = _make_sheet_rows(snapshot)
    state_row = snapshot["planner_state"][0]
    exported_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    metadata_rows = _metadata_rows(sheet_rows, exported_at, state_row["schema_version"], sample)
    wb = Workbook(write_only=True)
    wb.properties.title = "Financial Planner Data Template" if sample else "Financial Planner Data Backup"
    wb.properties.subject = "Complete application data in readable Excel worksheets"
    ws = wb.create_sheet("Metadata")
    _append_row(ws, HEADERS["Metadata"], header=True)
    for row in metadata_rows:
        _append_row(ws, row)
    for sheet in SHEETS[1:]:
        ws = wb.create_sheet(sheet)
        _append_row(ws, HEADERS[sheet], header=True)
        for row in sheet_rows[sheet]:
            _append_row(ws, row)
    buffer = BytesIO()
    wb.save(buffer)
    content = buffer.getvalue()
    if len(content) > (MAX_WORKBOOK_BYTES if max_workbook_bytes is None else max_workbook_bytes):
        raise WorkbookError("Workbook exceeds the 64 MiB file limit")
    return content, {"exported_at": exported_at, "format_version": FORMAT_VERSION,
                    "schema_version": state_row["schema_version"], "sample_data": sample}

def template_bytes(max_workbook_bytes=None):
    today = date.today().isoformat()
    sample = empty()
    sample["accounts"]["SAMPLE0001"] = {"id": "SAMPLE0001", "name": "Example bank (sample only)",
        "type": "bank", "currency": "SGD", "cpf_type": "", "archived": False}
    sample["events"].append({"id": "SAMPLE-EVENT-1", "kind": "opening_cash", "date": today,
        "order": 0, "status": "active", "actor": "sample",
        "data": {"account": "SAMPLE0001", "currency": "SGD", "amount": "100.00", "date": today}})
    snapshot = {"planner_state": [{"id": 1, "schema_version": 1, "data": sample}],
        "instrument_catalog": [], "portfolio_history": [], "stock_cache": []}
    return _build_xlsx(snapshot, sample=True, max_workbook_bytes=max_workbook_bytes)

def _parse_timestamp(value, field):
    if not isinstance(value, str):
        raise WorkbookError(f"{field} must be ISO 8601 text")
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise WorkbookError(f"{field} must be ISO 8601 text") from exc
    if parsed.tzinfo is None:
        raise WorkbookError(f"{field} must include a timezone")
    return parsed.astimezone(timezone.utc)

def _read_xlsx(content, max_upload_bytes=None):
    if not isinstance(content, bytes) or not content:
        raise WorkbookError("Workbook is empty")
    if len(content) > (MAX_UPLOAD_BYTES if max_upload_bytes is None else max_upload_bytes):
        raise WorkbookError("Workbook exceeds the 64 MiB upload limit")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > 128:
                raise WorkbookError("Workbook contains too many internal files")
            if sum(info.file_size for info in entries) > MAX_UNCOMPRESSED_BYTES:
                raise WorkbookError("Workbook exceeds the 200 MiB expanded size limit")
            forbidden = ("xl/vbaProject.bin", "xl/externalLinks/", "xl/embeddings/", "xl/activeX/",
                         "xl/connections.xml", "xl/queryTables/")
            if any(any(name.startswith(path) for path in forbidden) for name in names):
                raise WorkbookError("Macros, external links, embedded objects, and data connections are not allowed")
            for name in names:
                if name.endswith(".rels"):
                    try:
                        relationships = ElementTree.fromstring(archive.read(name))
                    except ElementTree.ParseError as exc:
                        raise WorkbookError("Workbook contains malformed internal relationships") from exc
                    if any(value.lower() == "external" for relation in relationships
                           for key, value in relation.attrib.items() if key.rsplit("}", 1)[-1].lower() == "targetmode"):
                        raise WorkbookError("External workbook relationships are not allowed")
    except (zipfile.BadZipFile, OSError, RuntimeError) as exc:
        raise WorkbookError("Upload is not a valid Excel .xlsx workbook") from exc
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise WorkbookError("Upload is not a readable Excel .xlsx workbook") from exc
    try:
        if tuple(wb.sheetnames) != SHEETS:
            raise WorkbookError("Workbook must contain the six required worksheets in order")
        if len(wb.defined_names):
            raise WorkbookError("Workbook defined names are not allowed")
        if any(wb[name].sheet_state != "visible" for name in SHEETS):
            raise WorkbookError("Hidden worksheets are not allowed")
        result = {}
        for sheet in SHEETS:
            ws = wb[sheet]
            iterator = ws.iter_rows(values_only=False)
            header = next(iterator, None)
            if not header or tuple(cell.value for cell in header) != HEADERS[sheet]:
                raise WorkbookError(f"{sheet} has missing or unexpected column headers")
            rows = []
            for cells in iterator:
                if any(cell.data_type == "f" for cell in cells):
                    raise WorkbookError("Formulas are not allowed in backup workbooks")
                values = tuple(cell.value for cell in cells)
                expected_columns = len(HEADERS[sheet])
                if len(values) > expected_columns and any(value is not None for value in values[expected_columns:]):
                    raise WorkbookError(f"{sheet} contains unexpected columns")
                values = values[:expected_columns] + (None,) * max(0, expected_columns - len(values))
                if all(value is None for value in values):
                    continue
                rows.append(values)
                if len(rows) + 1 > MAX_SHEET_ROWS:
                    raise WorkbookError(f"{sheet} exceeds Excel's worksheet row limit")
            result[sheet] = rows
    finally:
        wb.close()
    metadata, sheet_meta = {}, {}
    for row in result["Metadata"]:
        section, name, value, checksum, description = row
        if not isinstance(section, str) or not isinstance(name, str):
            raise WorkbookError("Metadata rows require text section and name values")
        key = (section, name)
        if key in metadata:
            raise WorkbookError("Metadata contains duplicate fields")
        metadata[key] = (value, checksum, description)
        if section == "sheet":
            if name in sheet_meta:
                raise WorkbookError("Metadata contains duplicate worksheet checksums")
            sheet_meta[name] = (value, checksum)
    required_meta = {("workbook", "workbook_type"), ("workbook", "workbook_role"),
        ("workbook", "format_version"), ("workbook", "app_version"), ("workbook", "schema_version"),
        ("workbook", "exported_at"), ("workbook", "sample_data")}
    if not required_meta <= metadata.keys():
        raise WorkbookError("Metadata is missing required workbook version fields")
    get = lambda name: metadata[("workbook", name)][0]
    if get("workbook_type") != WORKBOOK_TYPE or get("workbook_role") not in ("backup", "template"):
        raise WorkbookError("Unsupported workbook type")
    if get("format_version") != FORMAT_VERSION or get("schema_version") != 1 or get("app_version") != APP_VERSION:
        raise WorkbookError("Workbook or application schema version is incompatible")
    _parse_timestamp(get("exported_at"), "exported_at")
    if not isinstance(get("sample_data"), bool) or ((get("workbook_role") == "template") != get("sample_data")):
        raise WorkbookError("Workbook role and sample_data marker disagree")
    if set(sheet_meta) != set(SHEETS[1:]):
        raise WorkbookError("Metadata must contain one count and checksum for every data worksheet")
    for sheet in SHEETS[1:]:
        count, checksum = sheet_meta[sheet]
        if count != len(result[sheet]) or not isinstance(checksum, str):
            raise WorkbookError(f"Row count mismatch for {sheet}")
        if not secrets.compare_digest(checksum, _digest(result[sheet])):
            raise WorkbookError(f"Checksum mismatch for {sheet}")
    return metadata, result

def _field_map(sheet, rows):
    return [dict(zip(HEADERS[sheet], row)) for row in rows]

def _decode_nodes(rows, cache=False):
    by_record = {}
    for row in rows:
        if cache:
            record_id, node_id, parent_id, position, key, kind, value = row
        else:
            node_id, parent_id, position, key, kind, value = row
            record_id = "planner_state"
        if not isinstance(record_id, str) or not isinstance(node_id, str) or not node_id:
            raise WorkbookError("Node tables require text record and node IDs")
        if parent_id is not None and not isinstance(parent_id, str):
            raise WorkbookError("Parent Node ID values must be text")
        if isinstance(position, bool) or not isinstance(position, int) or position < 0:
            raise WorkbookError("Node Position values must be non-negative integers")
        if kind not in ("object", "list", "string", "integer", "number", "boolean", "null"):
            raise WorkbookError("Node Type contains an unknown value")
        record = by_record.setdefault(record_id, {})
        if node_id in record:
            raise WorkbookError("Node IDs must be unique within their record")
        record[node_id] = {"parent": parent_id, "position": position, "key": key, "type": kind, "value": value}
    decoded = {}
    for record_id, nodes in by_record.items():
        roots = [node_id for node_id, node in nodes.items() if node["parent"] is None]
        if len(roots) != 1:
            raise WorkbookError("Each node tree must have exactly one root")
        root = roots[0]
        if not cache and nodes[root]["type"] != "object":
            raise WorkbookError("Planner State root node must be an object")
        children = {}
        for node_id, node in nodes.items():
            parent = node["parent"]
            if parent is not None:
                if parent not in nodes or nodes[parent]["type"] not in ("object", "list"):
                    raise WorkbookError("Node parent references are invalid")
                children.setdefault(parent, []).append((node["position"], node["key"], node_id))
        for parent_id, entries in children.items():
            entries.sort(key=lambda item: item[0])
            if [entry[0] for entry in entries] != list(range(len(entries))):
                raise WorkbookError("Child positions must be unique and consecutive")
            if nodes[parent_id]["type"] == "object":
                keys = ["" if key is None else key for _, key, _ in entries]
                if any(not isinstance(key, str) for key in keys) or len(set(keys)) != len(keys):
                    raise WorkbookError("Object keys must be unique text values")
            elif any(key is not None for _, key, _ in entries):
                raise WorkbookError("List child rows must not contain object keys")
        visiting, built = set(), {}

        def build(node_id, depth=0):
            if depth > 256:
                raise WorkbookError("Node tree exceeds the supported nesting depth")
            if node_id in visiting:
                raise WorkbookError("Node tree contains a cycle")
            if node_id in built:
                return built[node_id]
            visiting.add(node_id)
            node, descendants = nodes[node_id], children.get(node_id, [])
            kind, value = node["type"], node["value"]
            if kind == "object":
                if value is not None:
                    raise WorkbookError("Object node values must be blank")
                output = {"" if key is None else key: build(child_id, depth + 1) for _, key, child_id in sorted(descendants)}
            elif kind == "list":
                if value is not None:
                    raise WorkbookError("List node values must be blank")
                output = [build(child_id, depth + 1) for _, _, child_id in sorted(descendants)]
            else:
                if descendants:
                    raise WorkbookError("Scalar nodes cannot have child nodes")
                if kind == "null":
                    if value is not None:
                        raise WorkbookError("Null nodes must have blank values")
                    output = None
                elif kind == "string":
                    output = "" if value is None else value
                    if not isinstance(output, str):
                        raise WorkbookError("String node value must be text")
                elif kind == "boolean":
                    if value not in ("true", "false"):
                        raise WorkbookError("Boolean node value must be true or false")
                    output = value == "true"
                elif kind == "integer":
                    if not isinstance(value, str):
                        raise WorkbookError("Integer node values must be exact text")
                    try:
                        output = int(value)
                    except ValueError as exc:
                        raise WorkbookError("Integer node value is invalid") from exc
                else:
                    if not isinstance(value, str):
                        raise WorkbookError("Number node values must be exact text")
                    try:
                        output = float(value)
                    except ValueError as exc:
                        raise WorkbookError("Number node value is invalid") from exc
                    if not math.isfinite(output) or repr(output) != value:
                        raise WorkbookError("Number node value is non-finite or not canonical")
            visiting.remove(node_id)
            built[node_id] = output
            return output

        decoded[record_id] = build(root)
        if len(built) != len(nodes):
            raise WorkbookError("Node table contains unreachable rows")
    return decoded

def _unique(rows, keys, table):
    indexed = [tuple(row.get(key) for key in keys) for row in rows]
    try:
        if len(set(indexed)) != len(indexed):
            raise WorkbookError(f"{table} contains duplicate keys")
    except TypeError as exc:
        raise WorkbookError(f"{table} contains invalid key values") from exc

def _validate_tables(metadata, rows_by_sheet):
    planner_trees = _decode_nodes(rows_by_sheet["Planner State"])
    if set(planner_trees) != {"planner_state"}:
        raise WorkbookError("Planner State must contain one complete tree")
    state_row = planner_trees["planner_state"]
    if not isinstance(state_row, dict) or set(state_row) != {"id", "schema_version", "data"} or state_row["id"] != 1 or state_row["schema_version"] != 1:
        raise WorkbookError("Planner State must reconstruct the supported planner_state row")
    data = state_row["data"]
    # Backups predating recurring schedules remain importable.
    if isinstance(data, dict):
        data.setdefault("recurring_schedules", [])
    expected = {"accounts": dict, "credit_accounts": dict, "events": list, "loans": dict,
                "prices": dict, "fx": dict, "receipts": dict, "recurring_schedules": list,
                "bot": dict, "provider_status": dict}
    if not isinstance(data, dict):
        raise WorkbookError("Planner State data must be an object")
    for key, kind in expected.items():
        if key not in data or not isinstance(data[key], kind):
            raise WorkbookError(f"Planner State is missing or has malformed {key}")
    bot = data["bot"]
    if isinstance(bot.get("offset", 0), bool) or not isinstance(bot.get("offset", 0), int) or bot.get("offset", 0) < 0:
        raise WorkbookError("Telegram bot offset must be a non-negative integer")
    if isinstance(data.get("revision"), bool) or not isinstance(data.get("revision"), int) or data["revision"] < 0:
        raise WorkbookError("Planner State revision must be a non-negative integer")
    accounts, credits, loans = data["accounts"], data["credit_accounts"], data["loans"]
    for account_id, account in accounts.items():
        if not isinstance(account, dict) or account.get("id") != account_id or account.get("type") not in ("bank", "brokerage", "cpf"):
            raise WorkbookError("Account IDs or types are invalid")
    for credit_id, credit in credits.items():
        if not isinstance(credit, dict) or credit.get("id") != credit_id or not isinstance(credit.get("cards"), dict):
            raise WorkbookError("Credit-card account IDs or cards are malformed")
        if any(not isinstance(card, dict) or card.get("id") != card_id for card_id, card in credit["cards"].items()):
            raise WorkbookError("Credit-card references are malformed")
    for loan_id, loan in loans.items():
        if not isinstance(loan, dict) or loan.get("id") != loan_id or not isinstance(loan.get("rates"), list):
            raise WorkbookError("Loan IDs or rates are malformed")
    schedule_ids = set()
    for schedule in data["recurring_schedules"]:
        if (not isinstance(schedule, dict) or not isinstance(schedule.get("id"), str)
                or schedule["id"] in schedule_ids or schedule.get("cadence") not in ("monthly", "annual")
                or schedule.get("status") not in ("active", "paused", "stopped")):
            raise WorkbookError("Recurring schedule IDs, cadence, or status are invalid")
        account = accounts.get(schedule.get("account"))
        day_number, month_number = schedule.get("day_of_month"), schedule.get("month_of_year")
        if (not account or account.get("type") != "bank"
                or not isinstance(day_number, int) or isinstance(day_number, bool)
                or not isinstance(schedule.get("amount"), str)
                or not isinstance(schedule.get("description"), str)
                or not isinstance(schedule.get("next_due_date"), str)):
            raise WorkbookError("Recurring schedule fields or bank account are invalid")
        if schedule["cadence"] == "monthly":
            if not 1 <= day_number <= 28 or month_number is not None:
                raise WorkbookError("Monthly schedule date is invalid")
        else:
            maximum = 28 if month_number == 2 else 30 if month_number in (4, 6, 9, 11) else 31
            if (not isinstance(month_number, int) or isinstance(month_number, bool)
                    or not 1 <= month_number <= 12 or not 1 <= day_number <= maximum):
                raise WorkbookError("Annual schedule date must be valid every year")
        try:
            amount = decimal(schedule["amount"], positive=True)
            decimal(schedule.get("initial_amount", schedule["amount"]), positive=True)
            next_due = date.fromisoformat(schedule["next_due_date"])
            if schedule.get("currency") != account.get("currency", "SGD"):
                raise ValueError
            if not schedule["description"].strip() or len(schedule["description"]) > 160:
                raise ValueError
            if schedule["cadence"] == "monthly" and next_due.day != day_number:
                raise ValueError
            if schedule["cadence"] == "annual" and (next_due.month != month_number or next_due.day != day_number):
                raise ValueError
            changes = schedule.get("amount_changes", [])
            if not isinstance(changes, list):
                raise ValueError
            previous_effective = None
            for change in changes:
                if not isinstance(change, dict) or not isinstance(change.get("effective_from_due_date"), str):
                    raise ValueError
                effective = date.fromisoformat(change["effective_from_due_date"])
                decimal(change.get("amount"), positive=True)
                if schedule["cadence"] == "monthly" and effective.day != day_number:
                    raise ValueError
                if schedule["cadence"] == "annual" and (effective.month != month_number or effective.day != day_number):
                    raise ValueError
                if previous_effective is not None and effective < previous_effective:
                    raise ValueError
                previous_effective = effective
        except (ValueError, TypeError) as exc:
            raise WorkbookError("Recurring schedule amount or next due date is invalid") from exc
        schedule_ids.add(schedule["id"])
    event_ids, orders = set(), set()
    for event in data["events"]:
        if not isinstance(event, dict) or not {"id", "order", "kind", "date", "status", "data"} <= event.keys():
            raise WorkbookError("Ledger event is missing required fields")
        if not isinstance(event["id"], str) or isinstance(event["order"], bool) or not isinstance(event["order"], int):
            raise WorkbookError("Ledger event IDs and order must be text and integers")
        if event["id"] in event_ids or event["order"] in orders:
            raise WorkbookError("Ledger events contain duplicate IDs or order values")
        event_ids.add(event["id"])
        orders.add(event["order"])
        if event["kind"] not in KINDS or event["status"] not in ("active", "void", "superseded") or not isinstance(event["data"], dict):
            raise WorkbookError("Ledger event kind, status, or data is invalid")
        try:
            if day(event["date"]) > date.today():
                raise WorkbookError("Ledger events cannot be dated in the future")
        except (Invalid, TypeError) as exc:
            raise WorkbookError("Ledger event date is invalid") from exc
    try:
        replay(data, date.today())
        credit_replay(data, date.today())
        for loan in loans.values():
            loan_balance(data, loan, date.today())
    except Exception as exc:
        raise WorkbookError(f"Ledger replay validation failed: {exc}") from exc

    catalog = _field_map("Instrument Catalog", rows_by_sheet["Instrument Catalog"])
    for row in catalog:
        if set(row) != set(HEADERS["Instrument Catalog"]) or not isinstance(row["active"], bool):
            raise WorkbookError("Instrument Catalog columns or active values are invalid")
        if row["exchange"] not in ("NYSE", "NASDAQ", "LSE", "SGX") or not isinstance(row["symbol"], str):
            raise WorkbookError("Instrument Catalog exchange or symbol is invalid")
        if row["asset_class"] not in (None, "equity", "etf"):
            raise WorkbookError("Instrument Catalog asset class is invalid")
        if row["currency"] is not None and (not isinstance(row["currency"], str) or len(row["currency"]) != 3):
            raise WorkbookError("Instrument Catalog currency is invalid")
        _parse_timestamp(row["updated_at"], "Instrument Catalog updated_at")
    _unique(catalog, ("exchange", "symbol"), "Instrument Catalog")

    history = _field_map("Portfolio History", rows_by_sheet["Portfolio History"])
    for row in history:
        if row["account"] not in accounts:
            raise WorkbookError("Portfolio History references an unknown account")
        try:
            date.fromisoformat(row["date"])
            _parse_timestamp(row["updated_at"], "Portfolio History updated_at")
            for field in ("value_sgd", "gross_change_sgd", "adjusted_change_sgd", "external_flow_sgd"):
                if row[field] is not None and not Decimal(str(row[field])).is_finite():
                    raise WorkbookError(f"Portfolio History {field} must be finite")
        except (ValueError, TypeError) as exc:
            raise WorkbookError("Portfolio History contains invalid dates or numeric values") from exc
    _unique(history, ("account", "date"), "Portfolio History")

    cache_meta, records = rows_by_sheet["Stock Cache"], {}
    for raw in cache_meta:
        record = dict(zip(HEADERS["Stock Cache"], raw))
        if not isinstance(record["Record ID"], str) or record["Record ID"] in records:
            raise WorkbookError("Stock Cache record IDs are invalid or duplicated")
        _parse_timestamp(record["fetched_at"], "Stock Cache fetched_at")
        _parse_timestamp(record["expires_at"], "Stock Cache expires_at")
        records[record["Record ID"]] = record
    _unique([{"security_id": row[1], "kind": row[2]} for row in cache_meta], ("security_id", "kind"), "Stock Cache")
    cache_payloads = _decode_nodes(rows_by_sheet["Stock Cache Data"], cache=True)
    if set(cache_payloads) != set(records):
        raise WorkbookError("Stock Cache and Stock Cache Data records do not match")

    return {"planner_state": state_row,
            "instrument_catalog": [{key: row[key] for key in HEADERS["Instrument Catalog"]} for row in catalog],
            "portfolio_history": [{key: row[key] for key in HEADERS["Portfolio History"]} for row in history],
            "stock_cache": [{"security_id": records[rid]["security_id"], "kind": records[rid]["kind"],
                "fetched_at": records[rid]["fetched_at"], "expires_at": records[rid]["expires_at"], "data": payload}
                for rid, payload in cache_payloads.items()]}

def _read_tables(content, max_upload_bytes=None):
    metadata, rows_by_sheet = _read_xlsx(content, max_upload_bytes=max_upload_bytes)
    tables = _validate_tables(metadata, rows_by_sheet)
    get = lambda name: metadata[("workbook", name)][0]
    return metadata, rows_by_sheet, tables, {"format_version": get("format_version"),
        "schema_version": get("schema_version"), "app_version": get("app_version"),
        "exported_at": get("exported_at"), "sample_data": get("sample_data")}

def read_workbook(content, max_upload_bytes=None):
    try:
        metadata, rows_by_sheet, tables, info = _read_tables(content, max_upload_bytes=max_upload_bytes)
    except WorkbookError:
        raise
    except (TypeError, ValueError, KeyError, IndexError, OverflowError, RecursionError) as exc:
        raise WorkbookError("Workbook contains invalid or unsupported worksheet values") from exc
    info["tables"] = tables
    info["counts"] = {sheet: len(rows_by_sheet[sheet]) for sheet in SHEETS[1:]}
    return info
