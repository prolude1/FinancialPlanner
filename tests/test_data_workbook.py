from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import zipfile

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, insert, select

from app.core import store
from app.core import data_workbook as workbook


@pytest.fixture
def isolated_db(tmp_path, monkeypatch):
    engine = create_engine("sqlite:///" + str(tmp_path / "workbook.db"))
    monkeypatch.setattr(store, "engine", engine)
    monkeypatch.setenv("DATA_ROLLBACK_DIR", str(tmp_path / "rollback"))
    workbook._confirmations.clear()
    store.migrate()
    yield engine
    engine.dispose()


def seed_tables():
    state = store.read()
    account_id = "0000000001"
    with store.transaction() as data:
        data["accounts"][account_id] = {"id": account_id, "name": "銀行 💵", "type": "bank",
            "currency": "SGD", "cpf_type": "", "archived": False, "native": {}, "label": "=literal"}
        data["events"].append({"id": "sample-event", "order": 1, "kind": "opening_cash",
            "date": "2026-01-02", "status": "active", "data": {"account": account_id, "amount": "12.340000"}})
        data["revision"] += 1
    now = datetime.now(timezone.utc)
    with store.engine.begin() as c:
        store.lock_application(c)
        c.execute(insert(store.instruments).values(exchange="NASDAQ", symbol="ABC", name="ABC Inc",
            asset_class="equity", currency="USD", native_exchange="NASDAQ", source="fixture",
            active=True, updated_at=now))
        c.execute(insert(store.stock_cache).values(security_id="0000000001", kind="overview",
            data={"decimal": "1.2300000000", "none": None, "false": False, "list": [], "formula": "=SUM(A1)", "empty": ""},
            fetched_at=now, expires_at=now))
        c.execute(insert(store.portfolio_history).values(account=account_id, date="2026-01-02", value_sgd="10.00",
            gross_change_sgd=None, adjusted_change_sgd=None, external_flow_sgd="10.00", updated_at=now))


def test_xlsx_round_trip_and_atomic_import_rollback(isolated_db):
    seed_tables()
    before = workbook.snapshot_bytes()[0]
    parsed = workbook.read_workbook(before)
    assert tuple(load_workbook(BytesIO(before), read_only=True).sheetnames) == workbook.SHEETS
    assert parsed["tables"]["planner_state"]["data"]["accounts"]["0000000001"]["label"] == "=literal"
    assert parsed["tables"]["stock_cache"][0]["data"] == {
        "decimal": "1.2300000000", "none": None, "false": False, "list": [], "formula": "=SUM(A1)", "empty": ""}
    preview = workbook.validate_workbook(before)
    with store.transaction() as data:
        data["accounts"]["0000000001"]["name"] = "Changed after export"
        data["revision"] += 1
        data["bot"]["offset"] = 9
    preview = workbook.validate_workbook(before)
    result = workbook.commit_import(preview["archive_sha256"], preview["current_revision"],
                                    preview["current_etag"], preview["confirmation_token"])
    assert result["status"] == "committed"
    restored = workbook.snapshot_bytes()[0]
    expected_tables = parsed["tables"]
    expected_tables["planner_state"]["data"]["bot"]["offset"] = 9
    assert workbook.read_workbook(restored)["tables"] == expected_tables
    assert store.read()["bot"]["offset"] == 9
    rollback = list(Path(isolated_db.url.database).parent.joinpath("rollback").glob("rollback-*.xlsx"))
    assert len(rollback) == 1
    assert workbook.read_workbook(rollback[0].read_bytes())["tables"]["planner_state"]["data"]["accounts"]["0000000001"]["name"] == "Changed after export"


def test_template_is_synthetic_and_import_preview_warns(isolated_db):
    sample, _ = workbook.template_bytes()
    parsed = workbook.read_workbook(sample)
    assert parsed["sample_data"] is True
    assert parsed["tables"]["planner_state"]["data"]["accounts"]
    preview = workbook.validate_workbook(sample)
    assert preview["sample_data"] and preview["importable"] and preview["confirmation_token"]
    assert any("synthetic sample data" in warning for warning in preview["warnings"])
    assert any("replaces current records" in warning for warning in preview["warnings"])


def test_tampering_formulas_and_external_relationships_are_rejected(isolated_db):
    seed_tables()
    content, _ = workbook.snapshot_bytes()
    modified = load_workbook(BytesIO(content))
    modified["Instrument Catalog"]["C2"] = "changed"
    tampered = BytesIO()
    modified.save(tampered)
    with pytest.raises(workbook.WorkbookError, match="Checksum mismatch"):
        workbook.read_workbook(tampered.getvalue())

    formula = load_workbook(BytesIO(content))
    formula["Instrument Catalog"]["C2"] = "=1+1"
    formula_stream = BytesIO()
    formula.save(formula_stream)
    with pytest.raises(workbook.WorkbookError, match="Formulas are not allowed"):
        workbook.read_workbook(formula_stream.getvalue())

    external = BytesIO()
    with zipfile.ZipFile(BytesIO(content)) as source, zipfile.ZipFile(external, "w") as target:
        for item in source.infolist():
            body = source.read(item.filename)
            if item.filename == "xl/_rels/workbook.xml.rels":
                body = body.replace(b"</Relationships>", b'<Relationship Id="rId99" Type="urn:external" Target="https://example.invalid" TargetMode="External"/></Relationships>')
            target.writestr(item, body)
    with pytest.raises(workbook.WorkbookError, match="External workbook relationships"):
        workbook.read_workbook(external.getvalue())


def test_upload_bounds_and_stale_confirmation(isolated_db, monkeypatch):
    content, _ = workbook.snapshot_bytes()
    monkeypatch.setattr(workbook, "MAX_UPLOAD_BYTES", 4)
    with pytest.raises(workbook.WorkbookError, match="64 MiB upload limit"):
        workbook.read_workbook(b"12345")
    monkeypatch.setattr(workbook, "MAX_UPLOAD_BYTES", 64 * 1024 * 1024)
    preview = workbook.validate_workbook(content)
    with store.transaction() as data:
        data["revision"] += 1
    with pytest.raises(workbook.WorkbookError, match="changed after validation"):
        workbook.commit_import(preview["archive_sha256"], preview["current_revision"],
                               preview["current_etag"], preview["confirmation_token"])
